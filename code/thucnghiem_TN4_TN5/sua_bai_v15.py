# -*- coding: utf-8 -*-
"""Sua Paper_Medical_PHT_v14.docx -> v15, THEO DUNG PHIEU TN4 va TN5.

Phieu TN4 muc 1: Hit@1 <= 5% -> "Add row vao Bang 12 -> tong 10 he thong.
Them bang chung cho Finding B/C". Vi vay:
  - THEM hang PhoGPT vao Bang 12 (thanh 10 he)
  - Sua moi cho ghi "bon he" / "Chin he thong" cho khop
  - Them mot doan neu ro han che cua so lieu nay (de thay quyet dinh, khong tu bo)
Phieu TN5 muc 5.3: them doan do nhay prompt.

Chay lai script nay tren v14 cho ra dung v15.
"""
import copy, hashlib, io, os, sys

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.run import Run

DO = RGBColor(0xFF, 0x00, 0x00)   # moi cho da sua deu to do de thay ra soat

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

NGUON = r'D:\Download\Paper_Medical_PHT_v14.docx'
DICH  = r'D:\Download\Paper_Medical_PHT_v15.docx'

# ---------------------------------------------------------------- thay chu
# (chuoi cu, chuoi moi, so lan phai thay)
THAY = [
    ("Bảng 12. Chín hệ thống trên 181 ca độc lập",
     "Bảng 12. Mười hệ thống trên 181 ca độc lập", 1),

    ("đồ thị; bốn hệ cuối tìm trực tiếp trên toàn bộ danh mục ICD-10 quốc gia",
     "đồ thị; năm hệ cuối tìm trực tiếp trên toàn bộ danh mục ICD-10 quốc gia", 1),

    ("chúng tôi chấm thêm bốn hệ không chịu ràng buộc ấy",
     "chúng tôi chấm thêm năm hệ không chịu ràng buộc ấy", 1),

    ("Chúng tôi vì thế chấm thêm bốn hệ không chịu ràng buộc đó",
     "Chúng tôi vì thế chấm thêm năm hệ không chịu ràng buộc đó", 1),

    ("Ba hệ còn lại là mô hình ngôn ngữ lớn sinh mã trực tiếp",
     "Bốn hệ còn lại là mô hình ngôn ngữ lớn sinh mã trực tiếp", 1),

    ("Phi-3.5-mini-instruct ở fp16, cùng giải mã tham lam và cùng giới hạn 200 token sinh.",
     "Phi-3.5-mini-instruct và PhoGPT-4B-Chat ở fp16, cùng giải mã tham lam, giới hạn "
     "200 token sinh với ba mô hình đầu và 250 token với PhoGPT-4B-Chat.", 1),

    ("BM25 tự do 8,29%, Qwen2.5 và Phi-3.5 cùng 7,73%. Cả bốn đều cao hơn từng hệ "
     "trong nhóm đóng danh mục, vốn nằm trong dải 1,66–6,08%.",
     "BM25 tự do 8,29%, Qwen2.5 và Phi-3.5 cùng 7,73%, còn PhoGPT-4B-Chat chỉ đạt 0,55%. "
     "Bốn hệ đầu đều cao hơn từng hệ trong nhóm đóng danh mục, vốn nằm trong dải "
     "1,66–6,08%; riêng PhoGPT-4B-Chat thấp hơn cả nhóm ấy.", 1),

    ("bốn hệ tự do đạt lần lượt 7,94% (Phi-3.5), 6,35% (BM25 tự do), 3,17% (Llama-3) "
     "và 3,17% (Qwen2.5)",
     "năm hệ tự do đạt lần lượt 7,94% (Phi-3.5), 6,35% (BM25 tự do), 3,17% (Llama-3), "
     "3,17% (Qwen2.5) và 0,00% (PhoGPT-4B-Chat)", 1),

    ("Hợp của bốn hệ tự do trên nhóm này cũng chỉ đạt 10/63 ca",
     "Hợp của năm hệ tự do trên nhóm này cũng chỉ đạt 10/63 ca", 1),

    # ---- doan ket luan (Muc 7): dong bo theo Bang 12 moi ----
    ("Ba thí nghiệm bổ sung ở Mục 6.5 vừa củng cố vừa giới hạn kết luận này.",
     "Bốn thí nghiệm bổ sung ở Mục 6.5 vừa củng cố vừa giới hạn kết luận này.", 1),

    ("Thứ ba, bốn hệ không bị ràng buộc bởi danh mục của đồ thị, tìm trên toàn bộ "
     "13.081 mã ICD-10 quốc gia, đạt 7,73–12,71% Hit@1 và 3,17–7,94% trên 63 ca nằm "
     "ngoài danh mục ấy, nơi năm hệ kia bằng không theo thiết kế.",
     "Thứ ba, năm hệ không bị ràng buộc bởi danh mục của đồ thị, tìm trên toàn bộ "
     "13.081 mã ICD-10 quốc gia, đạt 0,55–12,71% Hit@1 và 0,00–7,94% trên 63 ca nằm "
     "ngoài danh mục ấy, nơi năm hệ kia bằng không theo thiết kế. Thứ tư, kết quả của "
     "nhóm hệ tự do phụ thuộc cách viết prompt: chạy lại Llama-3-8B với ba biến thể "
     "prompt cho Hit@1 từ 9,39% đến 12,71%, và trên 118 ca với tới được, mức ý nghĩa so "
     "với cấu hình đề xuất chỉ vững dưới cách viết đã dùng ở Bảng 12 (p = 0,0041 với bản "
     "tiếng Việt, 0,0066 với bản tiếng Anh và 0,0768 với bản suy luận từng bước).", 1),

    # xuat hien 2 lan: tom tat dau bai (para 6) va ket luan (para 193).
    # So 29,83% (54/181) KHONG doi khi them PhoGPT - da tinh lai tu tep per-query.
    ("hợp của cả chín hệ chỉ đạt 29,83%",
     "hợp của cả mười hệ chỉ đạt 29,83%", 2),

    # ---- tom tat dau bai ----
    ("một hệ đối sánh từ vựng và ba mô hình ngôn ngữ lớn sinh mã tự do. Bốn hệ đạt "
     "7,73–12,71% Hit@1, và đạt 3,17–7,94% trên 63 ca nằm ngoài danh mục của đồ thị",
     "một hệ đối sánh từ vựng và bốn mô hình ngôn ngữ lớn sinh mã tự do. Năm hệ đạt "
     "0,55–12,71% Hit@1, và đạt 0,00–7,94% trên 63 ca nằm ngoài danh mục của đồ thị", 1),

    # ---- doan phan tich khoang cach (Muc 6.5) ----
    ("Con số cao nhất trong chín hệ là 12,71%",
     "Con số cao nhất trong mười hệ là 12,71%", 1),

    ("Hợp của cả chín hệ chỉ đạt 54/181 ca (29,83%)",
     "Hợp của cả mười hệ chỉ đạt 54/181 ca (29,83%)", 1),

    ("điều này giữ nguyên qua cả chín hệ. Thứ hai, một phần nhỏ quy được về độ phủ "
     "danh mục của đồ thị, đo bằng chênh lệch 0,00% so với 3,17–7,94% trên nhóm "
     "ngoài tầm với.",
     "điều này giữ nguyên qua cả mười hệ. Thứ hai, một phần nhỏ quy được về độ phủ "
     "danh mục của đồ thị, đo bằng chênh lệch 0,00% so với 0,00–7,94% trên nhóm "
     "ngoài tầm với.", 1),

    ("và phần này không hệ nào trong chín hệ thu hẹp được.",
     "và phần này không hệ nào trong mười hệ thu hẹp được.", 1),
]

# ------------------------------------------------------------ hang Bang 12
HANG_PHOGPT = ['PhoGPT-4B-Chat tự do (13.081 mã)', '0,55', '[0,10; 3,06]',
               '2,21', '0,0124', '0,85', '0,00']

# --------------------------------------------------------------- doan them
NEO = 'vốn nằm trong dải 1,66–6,08%'

DOAN_TN4 = (
    "Riêng hàng PhoGPT-4B-Chat cần đọc kèm một hạn chế về cách mô hình sinh đầu ra. "
    "Trên toàn bộ 181 ca, mô hình này chỉ sinh ra 13 mã ICD-10 ba ký tự khác nhau, trong "
    "đó mười mã trùng đúng danh sách minh hoạ trong prompt; khi thay danh sách minh hoạ "
    "bằng một bộ mã thuộc chuyên khoa khác thì đầu ra đổi theo bộ mã mới. Đầu ra vì vậy "
    "gần như không phụ thuộc mô tả bệnh nhân, và các chỉ số của hàng này phản ánh tần "
    "suất nhãn tham chiếu trùng với danh sách minh hoạ hơn là năng lực xếp hạng: một ca "
    "duy nhất trong 181 ca có nhãn ba ký tự trùng mã đứng đầu danh sách minh hoạ, đúng "
    "bằng Hit@1 = 0,55% đo được. Để loại trừ khả năng hiện tượng này do cách viết prompt, "
    "chúng tôi chạy lại toàn bộ 181 ca với một cách viết prompt thứ hai, khác cả về câu "
    "chữ lẫn cách bọc thẻ hội thoại; hai lượt cho văn bản sinh khác nhau ở 161/181 ca "
    "nhưng vẫn trùng nhau ở Hit@1 = 0,55%, ở cùng một ca đúng duy nhất, và ở đúng 13 mã "
    "ba ký tự, không cặp nào khác biệt theo kiểm định ghép cặp (b = 0, c = 0, p = 1). "
    "Kiểm định ghép cặp cũng cho thấy không ca nào PhoGPT xếp đúng mà cấu hình đề xuất "
    "xếp sai (b = 0), và mô hình này kém BM25 tự do một cách có ý nghĩa thống kê "
    "(b = 1 so với c = 15, McNemar p = 0,0005). Chúng tôi báo cáo hàng này để giữ tính "
    "đầy đủ của so sánh, đồng thời khuyến nghị không dùng nó làm bằng chứng về giới hạn "
    "của các mô hình bản ngữ tiếng Việt nói chung."
)

DOAN_TN5 = (
    "Vì mọi con số của nhóm hệ tự do đều đo dưới một cách viết prompt duy nhất, chúng tôi "
    "chạy lại Llama-3-8B trên cùng 181 ca với ba biến thể prompt gồm bản tiếng Việt gốc, "
    "bản tiếng Anh, và bản yêu cầu suy luận từng bước (chain-of-thought), giữ nguyên mọi "
    "tham số giải mã. Lượt chạy lại bản tiếng Việt tái lập đúng từng ca kết quả đã báo cáo "
    "ở trên (181/181 ca trùng nguyên văn đầu ra), xác nhận tính tất định của giải mã tham "
    "lam giữa hai phiên chạy khác nhau. Hit@1 trải từ 9,39% đến 12,71%, biên độ 3,32 điểm "
    "phần trăm: bản tiếng Việt 12,71%, bản tiếng Anh 11,60%, bản suy luận từng bước 9,39%. "
    "Chiều của chênh lệch không đổi dưới cả ba cách viết, khi số ca chỉ hệ tự do đúng luôn "
    "lớn hơn hẳn số ca chỉ cấu hình đề xuất đúng (20 so với 5, 19 so với 5, và 12 so với "
    "4). Tuy nhiên mức ý nghĩa thì phụ thuộc cách viết prompt: trên 118 ca với tới được, "
    "kiểm định McNemar cho p = 0,0041 với bản tiếng Việt, p = 0,0066 với bản tiếng Anh, và "
    "p = 0,0768 với bản suy luận từng bước, tức biến thể cuối không còn đạt ý nghĩa thống "
    "kê. Chúng tôi vì thế phát biểu kết luận ở mức bảo toàn: hệ tự do vượt cấu hình đề "
    "xuất trên nhóm ca mà đồ thị có nhãn, nhất quán về dấu dưới cả ba cách viết prompt, "
    "nhưng mức ý nghĩa chỉ vững dưới cách viết đã dùng ở Bảng 12. Đáng chú ý, biến thể suy "
    "luận từng bước hạ Hit@1 nhưng lại nâng Hit@5 lên 19,34%, cao nhất trong ba biến thể "
    "và cao hơn mức 14,36% của bản gốc, cho thấy chuỗi suy luận làm mô hình trải xác suất "
    "rộng hơn, đưa nhãn đúng vào top-5 thường xuyên hơn nhưng ít khi lên hạng nhất."
)


def _to_do(run):
    run.font.color.rgb = DO


def _chia_run(run, cu, moi, para):
    """Tach run thanh [truoc][moi - MAU DO][sau], giu nguyen dinh dang goc."""
    t = run.text
    i = t.index(cu)
    truoc, sau = t[:i], t[i + len(cu):]
    r_el = run._r

    def them_sau(neo_el, chu, do):
        el = copy.deepcopy(r_el)
        neo_el.addnext(el)
        r = Run(el, para)
        r.text = chu
        if do:
            _to_do(r)
        return el

    run.text = truoc
    el = them_sau(r_el, moi, True)
    if sau:
        them_sau(el, sau, False)
    if not truoc:
        r_el.getparent().remove(r_el)


def thay_trong_paragraph(p, cu, moi):
    """Thay chuoi va to do phan moi. Tra ve True neu co thay."""
    if cu not in p.text:
        return False
    for r in p.runs:                      # nam gon trong mot run
        if cu in r.text:
            _chia_run(r, cu, moi, p)
            return True
    # bi cat ngang nhieu run -> gop lai roi tach
    nguyen = p.text
    for r in p.runs[1:]:
        r.text = ''
    p.runs[0].text = nguyen
    _chia_run(p.runs[0], cu, moi, p)
    return True


def doan_moi_tu(neo_par, chu):
    p_moi = copy.deepcopy(neo_par._p)
    run_mau = None
    for con in list(p_moi):
        if con.tag == qn('w:pPr'):
            continue
        if con.tag == qn('w:r') and run_mau is None:
            run_mau = copy.deepcopy(con)
        p_moi.remove(con)
    if run_mau is None:
        run_mau = p_moi.makeelement(qn('w:r'), {})
    for con in list(run_mau):
        if con.tag != qn('w:rPr'):
            run_mau.remove(con)
    # to do ca doan moi
    rPr = run_mau.find(qn('w:rPr'))
    if rPr is None:
        rPr = run_mau.makeelement(qn('w:rPr'), {})
        run_mau.insert(0, rPr)
    for cu_color in rPr.findall(qn('w:color')):
        rPr.remove(cu_color)
    color = rPr.makeelement(qn('w:color'), {qn('w:val'): 'FF0000'})
    rPr.append(color)

    t = run_mau.makeelement(qn('w:t'), {})
    t.text = chu
    t.set(qn('xml:space'), 'preserve')
    run_mau.append(t)
    p_moi.append(run_mau)
    return p_moi


def main():
    d = Document(NGUON)

    # ---- 1. Thay chu ----
    print("=== Thay chu ===")
    for cu, moi, can in THAY:
        dem = 0
        for p in d.paragraphs:
            if thay_trong_paragraph(p, cu, moi):
                dem += 1
        for t in d.tables:                       # nhan de bang co the nam trong bang
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if thay_trong_paragraph(p, cu, moi):
                            dem += 1
        trang_thai = 'OK' if dem == can else f'*** {dem} lan, can {can} ***'
        print(f"  [{trang_thai}] {cu[:58]}...")
        assert dem == can, f"Khong thay dung so lan: {cu[:60]}"

    # ---- 2. Them hang PhoGPT vao Bang 12 ----
    print("\n=== Them hang vao Bang 12 ===")
    bang = None
    for t in d.tables:
        ten = [r.cells[0].text.strip() for r in t.rows]
        if any('Phi-3.5-mini tự do' in x for x in ten):
            bang = t
            break
    assert bang is not None, "Khong tim thay Bang 12"
    print(f"  Bang 12 truoc khi them: {len(bang.rows)} dong "
          f"({len(bang.rows)-1} he + 1 dong tieu de)")

    hang_cuoi = bang.rows[-1]
    hang_moi_el = copy.deepcopy(hang_cuoi._tr)
    hang_cuoi._tr.addnext(hang_moi_el)
    hang_moi = bang.rows[-1]
    for i, gia_tri in enumerate(HANG_PHOGPT):
        cell = hang_moi.cells[i]
        p = cell.paragraphs[0]
        if p.runs:
            p.runs[0].text = gia_tri
            _to_do(p.runs[0])
            for r in p.runs[1:]:
                r.text = ''
        else:
            _to_do(p.add_run(gia_tri))
        for extra in cell.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
    print(f"  Bang 12 sau khi them  : {len(bang.rows)} dong")
    print("  Hang moi:", ' | '.join(c.text for c in bang.rows[-1].cells))

    # ---- 3. Chen hai doan ----
    print("\n=== Chen hai doan vao muc 6.5 ===")
    neo = next(p for p in d.paragraphs if NEO in p.text)
    p_tn4 = doan_moi_tu(neo, DOAN_TN4)
    neo._p.addnext(p_tn4)
    p_tn5 = doan_moi_tu(neo, DOAN_TN5)
    p_tn4.addnext(p_tn5)
    print("  Da chen doan TN4 va doan TN5, style =", neo.style.name)

    d.save(DICH)
    print("\nDa luu:", DICH)

    # ---- 4. Kiem lai ----
    d2 = Document(DICH)
    ts = [p.text for p in d2.paragraphs]
    i = next(k for k, t in enumerate(ts) if NEO in t)
    print("\n=== KIEM LAI ===")
    print("  Paragraph :", len(ts), "(v14 co 276)")
    print("  Bang      :", len(d2.tables), "(v14 co 23)")
    b12 = next(t for t in d2.tables
               if any('PhoGPT' in r.cells[0].text for r in t.rows))
    print("  Bang 12   :", len(b12.rows) - 1, "he thong (can 10)")
    for r in b12.rows[1:]:
        print("     -", r.cells[0].text.strip(), "|", r.cells[1].text.strip())
    con_bon = [t for t in ts if 'bốn hệ tự do' in t or 'Chín hệ thống' in t
               or 'chấm thêm bốn hệ' in t]
    print("  Con sot 'bốn hệ tự do' / 'Chín hệ thống':",
          len(con_bon) if con_bon else "khong con")
    with open(DICH, 'rb') as f:
        print("  SHA-256 v15:", hashlib.sha256(f.read()).hexdigest())


if __name__ == '__main__':
    main()
