# -*- coding: utf-8 -*-
"""Sua Paper_Medical_PHT_v15.docx -> v16.

Ly do: luot TN5 chay dung phieu tung chu (§4.1: tokenizer tho, khong boc the hoi
thoai; V1 = 250 token) cho ket qua khac luot dau o bien the suy luan tung buoc.
So cua bai phai theo luot dung phieu.

  Bien the      Hit@1 A -> B      p(118) A -> B
  V1_baseline   12,71 -> 12,71    0,0041 -> 0,0041
  V2_english    11,60 -> 12,15    0,0066 -> 0,0043
  V3_cot         9,39 -> 14,36    0,0768 -> 0,0025   <-- lay lai y nghia thong ke

Nguon so: TN5B_KetQua.zip / tn5b_summary_variants.csv
Chay lai script nay tren v15 cho ra dung v16. Moi cho sua deu to do.
"""
import copy, hashlib, io, sys

from docx import Document
from docx.shared import RGBColor
from docx.text.run import Run

DO = RGBColor(0xFF, 0x00, 0x00)
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

NGUON = r'D:\codevstdio\repo\MedicalGraph\paper\Paper_Medical_PHT_v15.docx'
DICH  = r'D:\codevstdio\repo\MedicalGraph\paper\Paper_Medical_PHT_v16.docx'

MOC_TN5 = 'Vì mọi con số của nhóm hệ tự do đều đo dưới một cách viết prompt duy nhất'

DOAN_TN5_MOI = (
    "Vì mọi con số của nhóm hệ tự do đều đo dưới một cách viết prompt duy nhất, chúng tôi "
    "chạy lại Llama-3-8B trên cùng 181 ca với ba biến thể prompt gồm bản tiếng Việt gốc, "
    "bản tiếng Anh, và bản yêu cầu suy luận từng bước (chain-of-thought), giữ nguyên mọi "
    "tham số giải mã. Bản tiếng Việt tái lập đúng mức Hit@1 = 12,71% đã báo cáo ở Bảng 12 "
    "và đúng trên cùng 23 ca ấy, không lệch một ca nào. Hit@1 trải từ 12,15% đến 14,36%, "
    "biên độ 2,21 điểm phần trăm: bản tiếng Việt 12,71%, bản tiếng Anh 12,15%, bản suy "
    "luận từng bước 14,36%. Chiều của chênh lệch không đổi dưới cả ba cách viết, khi số ca "
    "chỉ hệ tự do đúng luôn lớn hơn hẳn số ca chỉ cấu hình đề xuất đúng (20 so với 5, 18 "
    "so với 4, và 21 so với 5), và mức ý nghĩa cũng giữ được ở cả ba: trên 118 ca với tới "
    "được, kiểm định McNemar cho p = 0,0041 với bản tiếng Việt, p = 0,0043 với bản tiếng "
    "Anh và p = 0,0025 với bản suy luận từng bước, cả ba vẫn đạt sau hiệu chỉnh Holm cho "
    "tám phép so sánh (p hiệu chỉnh lần lượt 0,0326, 0,0347 và 0,0200). Biến thể suy luận "
    "từng bước cho cả Hit@1 lẫn Hit@5 cao nhất, với Hit@5 = 22,65% so với 16,02% của bản "
    "gốc, nhưng cũng là biến thể duy nhất để lại ca không rút được mã nào từ văn bản sinh "
    "(4/181 ca, so với 0 ở hai biến thể kia), do chuỗi suy luận chiếm hết giới hạn token. "
    "Một lượt chạy song song, khác ở chỗ bọc đầu vào bằng thẻ hội thoại của mô hình thay "
    "vì đưa prompt thô, cho thấy độ nhạy này không chỉ nằm ở câu chữ: bản tiếng Việt vẫn "
    "cho đúng 12,71% và trùng nguyên văn từng ca với lượt ở Bảng 12, xác nhận tính tất "
    "định của giải mã tham lam giữa hai phiên chạy khác nhau, nhưng bản suy luận từng bước "
    "hạ xuống 9,39% và không còn đạt ý nghĩa thống kê (p = 0,0768). Chúng tôi vì vậy phát "
    "biểu: hệ tự do vượt cấu hình đề xuất trên nhóm ca mà đồ thị có nhãn, nhất quán cả về "
    "dấu lẫn mức ý nghĩa dưới cả ba cách viết prompt khi đầu vào được đưa theo đúng giao "
    "thức của thí nghiệm này, còn riêng với biến thể suy luận từng bước thì mức ý nghĩa "
    "còn phụ thuộc cách bọc đầu vào."
)

THAY = [
    # Sua mot cau co san tu v13, KHONG thuoc phieu nao. Cau nay bi sai o buoc
    # v12 -> v13: thong ke max theo tung he (= 1) bi dat vao mot cau co ngu phap
    # cua hop cac he (= 2). Tinh lai tu per-query: Qwen2.5 bat IND211 (Z20.3),
    # Phi-3.5 tu do bat IND163 (Z31.6); khong he nao bat qua 1 ca.
    # Truy nguon day du: NGUON_GOC_CHUONG_Z.md
    ("và 27 ca mang nhãn chương Z chỉ có đúng một ca được một hệ bắt trúng.",
     "và trong 27 ca mang nhãn chương Z, không hệ nào bắt trúng quá một ca, hợp của "
     "cả mười hệ cũng chỉ được hai ca.", 1),

    ("Thứ tư, kết quả của nhóm hệ tự do phụ thuộc cách viết prompt: chạy lại "
     "Llama-3-8B với ba biến thể prompt cho Hit@1 từ 9,39% đến 12,71%, và trên 118 ca "
     "với tới được, mức ý nghĩa so với cấu hình đề xuất chỉ vững dưới cách viết đã dùng "
     "ở Bảng 12 (p = 0,0041 với bản tiếng Việt, 0,0066 với bản tiếng Anh và 0,0768 với "
     "bản suy luận từng bước).",
     "Thứ tư, kết quả của nhóm hệ tự do có dao động theo cách viết prompt nhưng cả chiều "
     "lẫn mức ý nghĩa đều không đổi: chạy lại Llama-3-8B với ba biến thể prompt cho Hit@1 "
     "từ 12,15% đến 14,36%, và trên 118 ca với tới được, cả ba biến thể đều vượt cấu hình "
     "đề xuất có ý nghĩa thống kê (p = 0,0041 với bản tiếng Việt, 0,0043 với bản tiếng Anh "
     "và 0,0025 với bản suy luận từng bước, cả ba vẫn đạt sau hiệu chỉnh Holm cho tám phép "
     "so sánh).", 1),
]


def _to_do(run):
    run.font.color.rgb = DO


def _chia_run(run, cu, moi, para):
    t = run.text
    i = t.index(cu)
    truoc, sau = t[:i], t[i + len(cu):]
    r_el = run._r

    def them_sau(neo_el, chu, do):
        el = copy.deepcopy(r_el)
        neo_el.addnext(el)
        r = Run(el, para)
        r.text = chu
        if do:
            _to_do(r)
        return el

    run.text = truoc
    el = them_sau(r_el, moi, True)
    if sau:
        them_sau(el, sau, False)
    if not truoc:
        r_el.getparent().remove(r_el)


def thay_trong_paragraph(p, cu, moi):
    if cu not in p.text:
        return False
    for r in p.runs:
        if cu in r.text:
            _chia_run(r, cu, moi, p)
            return True
    nguyen = p.text
    for r in p.runs[1:]:
        r.text = ''
    p.runs[0].text = nguyen
    _chia_run(p.runs[0], cu, moi, p)
    return True


def main():
    d = Document(NGUON)

    print("=== 1. Viet lai doan TN5 o Muc 6.5 ===")
    ds = [p for p in d.paragraphs if p.text.startswith(MOC_TN5)]
    assert len(ds) == 1, f"Tim thay {len(ds)} doan TN5, can dung 1"
    p = ds[0]
    cu_len = len(p.text)
    p.runs[0].text = DOAN_TN5_MOI
    _to_do(p.runs[0])
    for r in p.runs[1:]:
        r.text = ''
    print(f"  {cu_len} -> {len(p.text)} ky tu, to do ca doan")

    print("\n=== 2. Sua menh de 'Thu tu' o phan ket luan ===")
    for cu, moi, can in THAY:
        dem = sum(1 for q in d.paragraphs if thay_trong_paragraph(q, cu, moi))
        print(f"  [{'OK' if dem == can else f'*** {dem}/{can} ***'}] {cu[:56]}...")
        assert dem == can, "Khong thay dung so lan"

    d.save(DICH)
    print("\nDa luu:", DICH)

    print("\n=== KIEM LAI ===")
    d2 = Document(DICH)
    ts = [q.text for q in d2.paragraphs]
    print("  Paragraph:", len(ts), "| Bang:", len(d2.tables),
          "| Anh:", len(d2.inline_shapes))
    b12 = next(t for t in d2.tables
               if any('PhoGPT' in r.cells[0].text for r in t.rows))
    print("  Bang 12  :", len(b12.rows) - 1, "he thong (can 10)")
    sot = {'0,0768': 0, '9,39% đến 12,71%': 0, 'biên độ 3,32': 0,
           '0,0066 với bản tiếng Anh': 0, '19,34%': 0, '11,60%, bản suy luận': 0,
           'chỉ có đúng một ca được một hệ bắt trúng': 0}
    can_co = {'0,0043': 0, '14,36%': 0, 'biên độ 2,21': 0, '0,0025': 0,
              '22,65%': 0, 'Holm': 0, 'không hệ nào bắt trúng quá một ca': 0}
    for t in ts:
        for k in sot:
            sot[k] += t.count(k)
        for k in can_co:
            can_co[k] += t.count(k)
    print("  So cu con sot :", {k: v for k, v in sot.items() if v})
    print("  So moi da co  :", can_co)
    with open(DICH, 'rb') as f:
        print("  SHA-256 v16:", hashlib.sha256(f.read()).hexdigest())


if __name__ == '__main__':
    main()
